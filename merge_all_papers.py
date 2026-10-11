#!/usr/bin/env python3
"""Merge every research-paper file into one Word document using docxcompose."""

import os
import shutil
import subprocess
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt, RGBColor
from docxcompose.composer import Composer

SRC = "/home/prem/Documents/research paper"
OUT_DOCX = os.path.join(SRC, "COMPLETE_MERGED_RISCV_RESEARCH_VOLUME.docx")
OUT_PDF = os.path.join(SRC, "COMPLETE_MERGED_RISCV_RESEARCH_VOLUME.pdf")

C_NAVY = RGBColor(16, 44, 87)
C_SLATE = RGBColor(53, 89, 143)
C_BODY = RGBColor(34, 34, 34)
C_GRAY = RGBColor(100, 110, 120)


def set_run(run, size=10, bold=False, color=C_BODY, italic=False):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def style_margins(doc):
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)


def heading(doc, text, size=16):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_run(r, size=size, bold=True, color=C_NAVY)


def body(doc, text, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run(r, size=size)


def bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("•  " + text)
    set_run(r, size=11)


def add_picture_centered(doc, path, width=Inches(6.4), caption=None):
    if not os.path.isfile(path):
        body(doc, f"[Missing file: {os.path.basename(path)}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(path, width=width)
    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(8)
        r = c.add_run(caption)
        set_run(r, size=9, italic=True, color=C_GRAY)


def pdf_to_pngs(pdf_path, dest_dir, prefix, dpi=130):
    out_prefix = os.path.join(dest_dir, prefix)
    subprocess.check_call(
        ["pdftoppm", "-png", "-r", str(dpi), pdf_path, out_prefix],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    return sorted(
        os.path.join(dest_dir, f)
        for f in os.listdir(dest_dir)
        if f.startswith(prefix) and f.endswith(".png")
    )


def make_cover(path):
    doc = Document()
    style_margins(doc)
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("COMPLETE RESEARCH VOLUME")
    set_run(r, size=13, bold=True, color=C_SLATE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run(
        "A Programmable Five-Stage RISC-V RV32I Microcontroller:\n"
        "From Assembly and Verilog to FPGA, Tiny Tapeout, and GDSII,\n"
        "and a Mission-Oriented Path Toward Satellite-on-Chip"
    )
    set_run(r, size=20, bold=True, color=C_NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    r = p.add_run("Prem Bahadur Rana  ·  Trishna Shrestha")
    set_run(r, size=13, bold=True, color=C_NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        "Nepal Space Foundation / Antarikchya Pratisthan Nepal, Kathmandu\n"
        "Open RTL: github.com/PrmBRana/ttgf_prem_programmable_RISC_v\n"
        "Tiny Tapeout GF 26a  ·  Project 0039  ·  4×4 tiles  ·  GF180MCU"
    )
    set_run(r, size=11, italic=True, color=C_GRAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    r = p.add_run("One Word file containing every paper, PDF, and figure from the research-paper folder.")
    set_run(r, size=11, color=C_BODY)

    heading(doc, "Contents of this merged file", 14)
    body(doc, "Nothing is left in a side file. Part I and Part II are the two full papers, "
              "including their tables, code listings, figures, and reference lists. "
              "The appendices paste the remaining PDFs and photographs.")
    bullet(doc, "Part I — From Verilog to GDSII (UART 16×, bootloader, GPIO, SPI modes, GTKWave, FPGA, 90 references)")
    bullet(doc, "Part II — Mission-oriented RISC-V RV32 microcontroller toward Satellite-on-Chip / NepaliSat-2")
    bullet(doc, "Appendix A — IAC 2026 RISC-V conference manuscript")
    bullet(doc, "Appendix B — Tiny Tapeout GF 26a invoice / receipt 61446")
    bullet(doc, "Appendix C — RISC-V instruction card")
    bullet(doc, "Appendix D — Extra figures (block diagram, lab photos, GTKWave, uploader, manuscript scans)")
    doc.save(path)


def make_divider(path, title, blurb):
    doc = Document()
    style_margins(doc)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(48)
    r = p.add_run(title)
    set_run(r, size=22, bold=True, color=C_NAVY)
    body(doc, blurb)
    doc.save(path)


def make_appendices(path, tmp):
    doc = Document()
    style_margins(doc)

    heading(doc, "APPENDIX A — IAC 2026 RISC-V Manuscript", 16)
    body(doc, "Full pages of IAC 2026_RISC_V (1).pdf.")
    pages = pdf_to_pngs(os.path.join(SRC, "IAC 2026_RISC_V (1).pdf"), tmp, "iac")
    for i, img in enumerate(pages, 1):
        add_picture_centered(doc, img, Inches(6.5), f"IAC 2026 manuscript, page {i} of {len(pages)}")

    heading(doc, "APPENDIX B — Tiny Tapeout GF 26a Invoice / Receipt 61446", 16)
    body(doc, "Official shuttle receipt for the 16-tile GF 26a allocation.")
    pages = pdf_to_pngs(os.path.join(SRC, "61446-en (1).pdf"), tmp, "invoice")
    for img in pages:
        add_picture_centered(doc, img, Inches(6.4), "Tiny Tapeout Ltd. Invoice / Receipt No. 61446")

    heading(doc, "APPENDIX C — RISC-V Instruction Card", 16)
    body(doc, "Reference card used while decoding the assembly listings.")
    pages = pdf_to_pngs(os.path.join(SRC, "RISCV_CARD.pdf"), tmp, "card")
    for i, img in enumerate(pages, 1):
        add_picture_centered(doc, img, Inches(6.5), f"RISC-V card, page {i} of {len(pages)}")

    heading(doc, "APPENDIX D — Extra Figures From the Research Folder", 16)
    body(doc, "Remaining images so the merged document holds the complete folder.")
    extras = [
        ("block_diagram_repo.png", "Repository block diagram from the GitHub RTL."),
        ("Untitled presentation.png", "Laboratory briefing architecture slide."),
        ("Screenshot from 2026-09-06 12-34-10.png", "Five-stage pipeline datapath drawing."),
        ("Screenshot from 2026-03-19 17-50-35.png", "GTKWave clock / waveform capture."),
        ("Screenshot from 2026-09-29 11-54-13.png", "UART Terminal & Instruction Uploader."),
        ("1789972494070.jpeg", "Vivado IP Integrator: Zynq UltraScale+, UART, bootloader."),
        ("1789972494046.jpeg", "Kria KV260 FPGA bench with FTDI USB–UART."),
        ("image (1).jpg", "Additional laboratory / presentation image."),
        ("page_code-09.png", "Earlier manuscript scan, page 09."),
        ("page_code-10.png", "Earlier manuscript scan, page 10."),
        ("page_code-11.png", "Earlier manuscript scan, page 11."),
        ("page_code-12.png", "Earlier manuscript scan, page 12."),
        ("page_code-13.png", "Earlier manuscript scan, page 13."),
    ]
    for name, cap in extras:
        add_picture_centered(doc, os.path.join(SRC, name), Inches(6.3), cap)

    heading(doc, "Folder inventory", 14)
    body(doc, "Word papers are in Parts I and II. Duplicate copies (paper.docx, paper_riscv_tutorial.docx, "
              "and the PDF twins of those two papers) are the same content, so they are not pasted a third time. "
              "IAC, invoice, and RISC-V card PDFs are Appendices A–C. "
              "Generator scripts remain in the folder as source code.")
    doc.save(path)


def main():
    tmp = tempfile.mkdtemp(prefix="riscv_merge_")
    cover = os.path.join(tmp, "cover.docx")
    part1d = os.path.join(tmp, "part1.docx")
    part2d = os.path.join(tmp, "part2.docx")
    apps = os.path.join(tmp, "apps.docx")

    make_cover(cover)
    make_divider(
        part1d,
        "PART I",
        "From Verilog to GDSII — the complete educational paper written from the taped-out RTL, "
        "including UART 16× oversampling, bootloader, GPIO, SPI modes, GTKWave, FPGA, Tiny Tapeout, "
        "and references [1]–[90]. The full paper begins on the next pages.",
    )
    make_divider(
        part2d,
        "PART II",
        "Mission-oriented RISC-V RV32 microcontroller as a first step toward Satellite-on-Chip "
        "for CubeSat applications — the complete SatOC paper, including assembly listings, "
        "FPGA results, and its reference list. The full paper begins on the next pages.",
    )
    make_appendices(apps, tmp)

    master = Document(cover)
    composer = Composer(master)
    composer.append(Document(part1d))
    composer.append(Document(os.path.join(SRC, "From_Verilog_to_GDSII_Programmable_RISCV_RV32I.docx")))
    composer.append(Document(part2d))
    composer.append(Document(os.path.join(SRC, "A_Mission_Oriented_RISCV_RV32_Microcontroller_Architecture_SatOC.docx")))
    composer.append(Document(apps))
    composer.save(OUT_DOCX)
    print("Wrote", OUT_DOCX, os.path.getsize(OUT_DOCX))

    subprocess.check_call(
        [
            "soffice", "--headless", "--convert-to", "pdf",
            "--outdir", SRC, OUT_DOCX,
        ]
    )
    print("Wrote", OUT_PDF, os.path.getsize(OUT_PDF) if os.path.isfile(OUT_PDF) else "missing")
    shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    main()
